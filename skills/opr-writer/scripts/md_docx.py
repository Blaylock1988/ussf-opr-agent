"""Export the working reports (_opr_work/ledger.md, research.md, record_review.md) as DOCX files in Output/.

Handles the Markdown the skill writes: headings, bullets, numbered items, pipe tables, **bold**, and
[links](url), which become "text (url)" so citations survive printing.
Controlled by draft.settings.export_reports (asked at intake; default true). opr.py build calls this.

Usage:  md_docx.py "<folder>/_opr_work" [--out "<folder>/Output"] [--name "<Name> OPR <period>"]
"""
import argparse
import re
from pathlib import Path

from docx import Document
from docx.shared import Mm, Pt

from common import load_json, utf8_stdout

REPORTS = {"ledger.md": "Fact Ledger", "research.md": "Research Report", "record_review.md": "Record Review"}
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def _runs(p, text, size=10):
    text = LINK.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text).replace("`", "")
    for n, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            r = p.add_run(part)
            r.bold = n % 2 == 1
            r.font.size = Pt(size)


def md_to_docx(md_path, out_path, title):
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Mm(15)
    doc.add_paragraph("CUI (When filled out)").alignment = 1
    lines = Path(md_path).read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                cols = max(len(r) for r in rows)
                table = doc.add_table(rows=0, cols=cols)
                table.style = "Table Grid"
                for n, cells in enumerate(rows):
                    for cell, text in zip(table.add_row().cells, cells + [""] * (cols - len(cells))):
                        _runs(cell.paragraphs[0], f"**{text}**" if n == 0 and text else text, 9)
            continue
        m = re.match(r"(#{1,4})\s+(.*)", ln)
        if m:
            doc.add_heading(LINK.sub(r"\1", m.group(2)), level=min(len(m.group(1)), 3))
        elif re.match(r"\s*[-*]\s+", ln):
            _runs(doc.add_paragraph(style="List Bullet" if not ln.startswith("  ") else "List Bullet 2"),
                  re.sub(r"^\s*[-*]\s+", "", ln))
        elif re.match(r"\s*\d+\.\s+", ln):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\.\s+", "", ln))
        elif ln.strip() and not ln.startswith("```"):
            _runs(doc.add_paragraph(), ln)
        i += 1
    doc.core_properties.title = title
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def export_reports(work_dir, out_dir=None, name=None):
    """Write each report that exists; return the paths written."""
    work = Path(work_dir)
    draft = load_json(work / "draft.json") if (work / "draft.json").is_file() else {}
    if not name:
        ratee = draft.get("ratee", {})
        name = f"{(ratee.get('name') or 'OPR').split(',')[0].strip()} OPR {ratee.get('period_thru', '')}".strip()
    out_dir = Path(out_dir) if out_dir else work.parent / "Output"
    written = []
    for fname, title in REPORTS.items():
        src = work / fname
        if src.is_file():
            written.append(md_to_docx(src, out_dir / f"{name} - {title}.docx", title))
    return written


def main():
    utf8_stdout()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("work_dir")
    ap.add_argument("--out")
    ap.add_argument("--name")
    args = ap.parse_args()
    for p in export_reports(args.work_dir, args.out, args.name) or ["(no ledger.md, research.md or record_review.md found)"]:
        print(f"wrote {p}" if isinstance(p, Path) else p)


if __name__ == "__main__":
    main()
