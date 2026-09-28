"""Extract every input file in a user's OPR folder to text, classify it, and index large files.

Writes into <folder>/_opr_work/:
  inputs/<name>.txt           plain text of each file (Unicode spaces preserved)
  inputs/<name>.chunks/       chunked copy of any file over ~25K tokens (never read these whole)
  manifest.json               kind, size, token estimate, notes and suggested follow-up per file
  history.json                prior OPR lines split by year/section (from history files)
  approved/afpc_acronyms.json parsed AFPC acronym page, when one is present

Usage:  extract_inputs.py "<folder>"
"""
import argparse
import csv
import html
import re
from pathlib import Path

from common import normalize_spaces, save_json, utf8_stdout, write_text

LARGE_TOKENS = 25_000
CHUNK_LINES = 300
SKIP_DIRS = {"_opr_work", "Output", ".git", "node_modules"}
TEXT_EXT = {".txt", ".md"}


def est_tokens(text):
    return len(text) // 4


def read_docx(path):
    from docx import Document

    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def read_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
            reader.pages[0].extract_text()
        except Exception:
            return ("[encrypted PDF - text not extractable here. If this is the official fillable DAF 707 (XFA), "
                    "use a 'print to PDF' copy; the form layout is already in the bundled form profile.]")
    pages = []
    for i, page in enumerate(reader.pages):
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # malformed page
            pages.append(f"[page {i + 1} unreadable: {exc}]")
    return "\n\f".join(pages)


def read_xlsx(path):
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).strip() for v in row]
            if any(cells):
                out.append("\t".join(cells).rstrip())
    return "\n".join(out)


def read_csv(path):
    with open(path, encoding="utf-8-sig", errors="replace") as f:
        return "\n".join("\t".join(r) for r in csv.reader(f))


def read_html(path):
    raw = Path(path).read_text(encoding="utf-8", errors="replace")
    raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
    raw = re.sub(r"(?i)</(tr|p|div|li|h\d|br)>|<br\s*/?>", "\n", raw)
    raw = re.sub(r"(?i)</t[dh]>", "\t", raw)
    return html.unescape(re.sub(r"<[^>]+>", "", raw))


READERS = {".docx": read_docx, ".pdf": read_pdf, ".xlsx": read_xlsx, ".xlsm": read_xlsx, ".csv": read_csv,
           ".html": read_html, ".htm": read_html}


def classify(path, text, root):
    rel = str(path.relative_to(root)).lower()
    name = path.name.lower()
    if "33-337" in name or ("tongue and quill" in text[:600].lower() and len(text) > 400_000):
        return "guidance:tongue_and_quill"
    if re.search(r"\b(daf|af)\s*(form\s*)?707\b|daf707", name):
        return "guidance:form"
    if ("afpc" in name or "air force acronym" in name) and "acronym" in name:
        return "guidance:afpc_acronyms"
    if "selected works" in name or "selected_works" in name or re.search(r"c-?note", name):
        return "guidance:cso_works"
    if "competitive categor" in name or "developmental categor" in name:
        return "guidance:competitive_categories"
    if "career review" in name or "assessment worksheet" in name:
        return "guidance:career_review"
    if "verb" in name or "game changer" in name:
        return "guidance:verbs"
    if re.search(r"acronym.*abbreviation.*list|glossary", name):
        return "guidance:glossary"
    if "attachment" in name and "ACRONYMS" in text and "ABBREVIATIONS" in text:
        return "guidance:approved_lists"
    if re.search(r"(previous|prior|past|history|old).*(opr|bullets)|(opr|bullets).*(previous|prior|past|history)", name) \
            or text.count("Rater Overall Assessment") >= 1:
        return "history"
    if "guid" in rel or "writing guide" in name or "handbook" in name:
        return "guidance"
    return "accomplishments"


def chunk(text, out_dir):
    out_dir.mkdir(parents=True, exist_ok=True)
    lines = text.splitlines()
    index = []
    for n, start in enumerate(range(0, len(lines), CHUNK_LINES)):
        seg = lines[start:start + CHUNK_LINES]
        heading = next((ln.strip() for ln in seg if re.match(r"^\s*(chapter|section|attachment|\d+\.\d+\.?)\s", ln, re.I)), seg[0].strip() if seg else "")
        f = out_dir / f"{n:03d}.txt"
        write_text(f, "\n".join(seg))
        index.append({"file": f.name, "lines": [start + 1, start + len(seg)], "first_heading": heading[:120]})
    save_json(out_dir / "index.json", index)
    return len(index)


def parse_history(text):
    """Split prior-OPR text into lines tagged by year and section; unmerge lines joined on extraction."""
    out, year, section = [], None, None
    for raw in text.splitlines():
        ln = normalize_spaces(raw).strip()
        if re.fullmatch(r"(19|20)\d\d", ln):
            year = ln
            continue
        low = ln.lower()
        if low.startswith("job description"):
            section = "job_description"
        elif low.startswith("additional rater"):
            section = "additional_rater"
        elif low.startswith("rater overall") or low.startswith("rater's overall"):
            section = "rater"
        elif low.startswith("reviewer"):
            section = "reviewer"
        elif low.startswith("acronyms"):
            section = "acronyms"
        elif ln.startswith("- ") and section and section != "acronyms":
            for part in re.split(r"\s-\s(?=[#A-Z])", ln[2:]):
                out.append({"year": year, "section": section, "text": "- " + part.strip()})
        elif section == "acronyms" and ln:
            out.append({"year": year, "section": "acronyms", "text": ln})
    return out


def parse_afpc(text):
    entries = {}
    for ln in text.splitlines():
        m = re.match(r"^\s*([A-Z][A-Za-z0-9&/\-]{1,15})\s*(?:\t|\s[-:|]\s|\s{2,})\s*(.{3,120})$", ln)
        if m and not m.group(2).isupper():
            entries.setdefault(m.group(1), [])
            if m.group(2).strip() not in entries[m.group(1)]:
                entries[m.group(1)].append(m.group(2).strip())
    return entries


def main():
    utf8_stdout()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder")
    args = ap.parse_args()
    root = Path(args.folder).resolve()
    work = root / "_opr_work"
    inputs = work / "inputs"
    manifest, history = [], []
    kinds_seen = set()

    for path in sorted(p for p in root.rglob("*") if p.is_file()):
        if any(part in SKIP_DIRS for part in path.relative_to(root).parts) or path.name.startswith("~$"):
            continue
        ext = path.suffix.lower()
        if ext in READERS:
            reader = READERS[ext]
        elif ext in TEXT_EXT:
            reader = lambda p: Path(p).read_text(encoding="utf-8", errors="replace")  # noqa: E731
        else:
            reader = None
        entry = {"path": str(path.relative_to(root)), "ext": ext}
        if reader is None:
            entry.update({"kind": "skipped", "note": "unsupported file type"})
            manifest.append(entry)
            continue
        try:
            text = reader(path)
        except Exception as exc:
            entry.update({"kind": "error", "note": f"could not read: {exc}"})
            manifest.append(entry)
            continue
        kind = classify(path, text, root)
        kinds_seen.add(kind)
        out = inputs / (path.relative_to(root).as_posix().replace("/", "__") + ".txt")
        write_text(out, text)
        tokens = est_tokens(text)
        entry.update({"kind": kind, "text_file": str(out.relative_to(work)), "chars": len(text), "est_tokens": tokens})
        notes = []
        if "please wait" in text[:400].lower() and "adobe" in text[:600].lower():
            notes.append("Dynamic XFA form: text not extractable. Use a 'print to PDF' copy or the Read tool for layout.")
        if ext == ".pdf" and len(text.strip()) < 200 and "encrypted" not in text:
            notes.append("Image-only PDF (no text layer): view it with the Read tool and transcribe what is needed.")
        if kind == "guidance:cso_works":
            notes.append("CSO works: bundled index covers Saltzman's Selected Works (C-Notes #1-#42). For a newer "
                         "compilation run build_cso_index.py <pdf> and search with ref_lookup.py --source cso.")
        if kind == "guidance:competitive_categories":
            notes.append("Competitive/developmental categories are bundled in data/competitive_categories.json; "
                         "update it if this table is newer. audience.py uses it to set the board audience.")
        if kind == "guidance:career_review":
            notes.append("Career Review worksheet: run career_review.py on _opr_work to score strats/pushes like a board.")
        if tokens > LARGE_TOKENS:
            n = chunk(text, out.with_suffix(".chunks"))
            entry["chunks"] = str(out.with_suffix(".chunks").relative_to(work))
            notes.append(f"LARGE (~{tokens:,} tokens): never read whole; {n} chunks + index.json")
        if kind == "guidance:tongue_and_quill":
            ed = re.search(r"(\d{1,2} [A-Z]{3,9} \d{4})", text[:2000])
            entry["edition"] = ed.group(1) if ed else "unknown"
            notes.append("Tongue and Quill: use the bundled digest / ref_lookup.py; rebuild the index with build_tq_index.py only for a newer edition.")
        if kind == "guidance:approved_lists":
            notes.append("Command acronym/abbreviation lists: if not the bundled CFC 26Jun2026 lists, run "
                         "build_approved_lists.py <pdf> --out _opr_work/approved to override.")
        if kind == "guidance:afpc_acronyms":
            afpc = parse_afpc(text)
            if not afpc:
                notes.append("The bundled reference/data/afpc_acronyms.json already holds the AFPC list (28 Oct 24 "
                             "update). Re-transcribe only if this export is newer.")
            if afpc:
                save_json(work / "approved" / "afpc_acronyms.json", {"source": path.name, "entries": afpc})
                notes.append(f"parsed {len(afpc)} AFPC acronyms into approved/afpc_acronyms.json (spot-check it)")
        if kind in ("guidance:verbs", "guidance:glossary"):
            notes.append("Rebuild lexicons with build_lexicons.py to merge this file.")
        if kind == "history":
            h = parse_history(text)
            history += h
            notes.append(f"{sum(1 for x in h if x['section'] != 'acronyms')} prior OPR lines parsed")
        entry["notes"] = notes
        manifest.append(entry)

    warnings = []
    if not any(k.startswith("guidance") for k in kinds_seen):
        warnings.append("No guidance found; bundled CFC lists and SpOC Sep 2023 body rules will be used.")
    if "guidance:approved_lists" in kinds_seen and not any(
            m.get("kind") == "guidance" and "writing guide" in m["path"].lower() and "attachment" not in m["path"].lower()
            for m in manifest):
        warnings.append("Only writing-guide ATTACHMENTS found, not the guide body; body rules default to the SpOC "
                        "Writing Guide (Sep 2023). Add the current command writing guide if available.")
    if not history:
        warnings.append("No prior OPRs found; style profile will use defaults.")
    save_json(work / "manifest.json", {"root": str(root), "files": manifest, "warnings": warnings})
    save_json(work / "history.json", {"lines": history})

    for m in manifest:
        extra = f" ~{m['est_tokens']:,} tok" if "est_tokens" in m else ""
        print(f"{m['kind']:28} {m['path']}{extra}")
        for n in m.get("notes", []):
            print(f"{'':28}   - {n}")
    for w in warnings:
        print("WARNING: " + w)
    print(f"\nmanifest: {work / 'manifest.json'}")


if __name__ == "__main__":
    main()
