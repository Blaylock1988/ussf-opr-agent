"""Render a draft to a DOCX that mirrors the DAF 707 layout, plus a rater-notes appendix.

Page 1 mirrors Secs I, II, IV, V, (VI) and X in Times New Roman 12 pt with text boxes at the
form's line width, keeping the fitted Unicode spaces so each line pastes flush into the form.
Each Box I field, the duty title and the last-feedback date sit on their own table rows. The fitness
score goes in Box I by default (ratee.fitness_placement: box1 | remarks | both | none).
Appendix: notes for the rater (pending strat/push decisions, alternates, ledger sources), fit table,
4 spare bullets, and lint findings.

Usage:  opr_docx.py draft.json [--out "<folder>/Output/<Name> OPR <period>.docx"]
"""
import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from common import iter_lines, load_form, load_json, normalize_spaces, utf8_stdout
from opr_acronyms import build_remarks
from rules import Approvals

PLACEHOLDER = re.compile(r"(#\[N\]/\[M\]|(?<!\w)\[[^\[\]\n]+\])")  # not line refs like rater[4]
WRAP_GUARD_MM = 0.3  # a hair of slack so Word never wraps a line fitted to the exact edge


def _font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold, run.italic = bold, italic
    return run


def _para(container, text="", size=12, bold=False, italic=False, highlight_placeholders=False):
    p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    if highlight_placeholders:
        for part in PLACEHOLDER.split(text):
            if part:
                r = _font(p.add_run(part), size, bold, italic)
                if PLACEHOLDER.fullmatch(part):
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    elif text:
        _font(p.add_run(text), size, bold, italic)
    return p


def _zero_cell_margins(table):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), "0" if side in ("left", "right") else "20")
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def _box(doc, width_mm, header, *rows):
    """A bordered one-column box sized to the form's line width: a header row, then one table row per
    entry in rows. Each entry is a list of lines, or a single string for a one-line row (e.g. DUTY TITLE)."""
    table = doc.add_table(rows=1 + len(rows), cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _zero_cell_margins(table)
    for row in table.rows:
        row.cells[0].width = Mm(width_mm + WRAP_GUARD_MM)
    hdr = table.rows[0].cells[0]
    hdr.paragraphs[0].text = ""
    _font(hdr.paragraphs[0].add_run(header), 9, bold=True)
    for row, lines in zip(table.rows[1:], rows):
        body = row.cells[0]
        body.paragraphs[0].text = ""
        for n, line in enumerate([lines] if isinstance(lines, str) else lines):
            p = body.paragraphs[0] if n == 0 else body.add_paragraph()
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
            for part in PLACEHOLDER.split(line):
                if part:
                    r = _font(p.add_run(part))
                    if PLACEHOLDER.fullmatch(part):
                        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def _fitness_placement(ratee, form):
    """box1 (default) | remarks | both | none. The DAF 707 has no fitness block yet, so Box I is the default."""
    fitness = form.get("optional_blocks", {}).get("fitness", {})
    if not ratee.get("fitness_score") or not fitness.get("enabled", True):
        return set(), None
    place = ratee.get("fitness_placement") or fitness.get("placement", "box1")
    where = {"box1": {"box1"}, "remarks": {"remarks"}, "both": {"box1", "remarks"}}.get(place, set())
    return where, fitness.get("template", "FITNESS SCORE: {score}").format(score=ratee["fitness_score"])


# Line and Width / status are narrow so the sources and alternates columns get the room (share of page width).
LINE_TABLE_MM = (0.125, 0.125, 0.36, 0.39)


def _table(doc, width_mm, headers, shares, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    widths = [Mm(width_mm * s) for s in shares]
    for col, w in zip(table.columns, widths):  # sets tblGrid/gridCol, which Word's fixed layout follows
        col.width = w
    for c, h, w in zip(table.rows[0].cells, headers, widths):
        c.width = w
        c.paragraphs[0].text = ""
        _font(c.paragraphs[0].add_run(h), 9, bold=True)
    for vals in rows:
        for cell, v, w in zip(table.add_row().cells, vals, widths):
            cell.width = w
            cell.paragraphs[0].text = ""
            _font(cell.paragraphs[0].add_run(v), 8)
    return table


def render(draft_path, out_path=None, lint_findings=None):
    path = Path(draft_path)
    draft = load_json(path)
    form = load_form(draft.get("form"), path.parent)
    width = form["line_width_mm"]
    ratee = draft.get("ratee", {})
    secs = draft.get("sections", {})

    doc = Document()
    sect = doc.sections[0]
    sect.page_width, sect.page_height = Mm(215.9), Mm(279.4)
    side = max((215.9 - width - WRAP_GUARD_MM) / 2 - 0.5, 5)
    sect.left_margin = sect.right_margin = Mm(side)
    sect.top_margin = sect.bottom_margin = Mm(10)

    _para(doc, "CUI (When filled out)", 8).alignment = 1
    _para(doc, "OFFICER PERFORMANCE REPORT (Lt thru Col)", 12, bold=True).alignment = 1
    _para(doc, f"{form['id'].upper()} working copy - paste each line into the official form", 8, italic=True).alignment = 1

    fit_where, fit_line = _fitness_placement(ratee, form)
    # one table row per identification field, so each pastes into its own myEval field
    ident = [f"{k.replace('_', ' ').upper()}: {ratee[k]}" for k in form.get("identification_fields", []) if ratee.get(k)]
    if "box1" in fit_where:
        ident.append(fit_line)
    _box(doc, width, "I. RATEE IDENTIFICATION DATA", *(ident or ["[complete in myEval]"]))

    for sec in form["sections"]:
        data = secs.get(sec["key"], {})
        lines = [ln.get("fitted") or ln["text"] for ln in data.get("lines", [])]
        if not lines and not sec.get("required"):
            continue
        header = f"{sec['label']} (Limit text to {sec['max_lines']} lines)"
        rows = [lines or ["[empty]"]]
        if sec.get("has_duty_title"):
            rows.insert(0, f"DUTY TITLE: {data.get('duty_title', '[duty title]')}")
        if sec["key"] == "rater":
            rows.append(f"Last performance feedback was accomplished on: {ratee.get('last_feedback_date', '[date]')}")
        _box(doc, width, header, *rows)

    remarks = draft.get("remarks")
    if remarks is None:
        remarks, _, _ = build_remarks(draft, form, Approvals(path.parent))
    remark_lines = [remarks or "(no acronyms require definition)"]
    if "remarks" in fit_where:
        remark_lines.append(fit_line)
    _box(doc, width, form["remarks"]["label"], remark_lines)

    # ------------------------------------------------------------ appendix
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _para(doc, "APPENDIX - NOTES FOR THE RATING CHAIN (not part of the OPR)", 12, bold=True)
    pending = [f"{sec_key}[{i + 1}]: {normalize_spaces(line['text'])}"
               for sec_key, i, line in iter_lines(draft, form) if PLACEHOLDER.search(line["text"])]
    _para(doc, "Decisions for the rater / additional rater", 11, bold=True)
    for item in pending or ["None - all stratifications and pushes are filled in."]:
        _para(doc, "- " + item, 10, highlight_placeholders=True)
    for key, opts in draft.get("push_options", {}).items():
        _para(doc, f"Push options for {key}: " + " | ".join(opts), 10)

    ledger = {e["id"]: e for e in draft.get("ledger", [])}

    def sources(line):
        return "; ".join(f"{lid}: {ledger[lid]['summary']} ({ledger[lid].get('source', '')})" if lid in ledger else lid
                         for lid in line.get("ledger", []))

    _para(doc, "Line-by-line sources, fit and alternates", 11, bold=True)
    rows = [(f"{k}[{i + 1}]", line) for k, i, line in iter_lines(draft, form)]
    _table(doc, width, ("Line", "Width / status", "Backed by (ledger)", "Alternates"), LINE_TABLE_MM,
           [(where, f"{line.get('width_mm', '?')} mm {line.get('status', 'unfitted')}", sources(line),
             "\n".join(normalize_spaces(a["text"] if isinstance(a, dict) else a) for a in line.get("alternates", [])))
            for where, line in rows])

    spares = draft.get("spares", [])
    _para(doc, "Spare performance bullets (fitted, lint-clean swap-ins)", 11, bold=True)
    if spares:
        for sp in spares:
            _para(doc, sp.get("fitted") or sp["text"], 12, highlight_placeholders=True)
        _table(doc, width, ("Spare", "Width / status", "Backed by (ledger)", "Swap for"), LINE_TABLE_MM,
               [(f"spares[{n + 1}]", f"{sp.get('width_mm', '?')} mm {sp.get('status', 'unfitted')}", sources(sp),
                 sp.get("swap_for", "any performance line"))
                for n, sp in enumerate(spares)])
    else:
        _para(doc, "None drafted yet.", 10)

    if lint_findings:
        _para(doc, "Open lint findings", 11, bold=True)
        for f in lint_findings:
            _para(doc, f"{f['severity'].upper()} {f['where']}: {f['message']}", 9)

    if not out_path:
        name = (ratee.get("name") or "OPR").split(",")[0].strip()
        period = ratee.get("period_thru", "")
        out_path = path.parent.parent / "Output" / f"{name} OPR {period}".strip()
        out_path = Path(str(out_path) + ".docx")
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def main():
    utf8_stdout()
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("draft")
    ap.add_argument("--out")
    ap.add_argument("--no-lint", action="store_true")
    args = ap.parse_args()
    findings = None
    if not args.no_lint:
        from opr_lint import lint

        findings = [f for f in lint(args.draft) if f["severity"] in ("error", "warning", "judgment")]
    out = render(args.draft, args.out, findings)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
